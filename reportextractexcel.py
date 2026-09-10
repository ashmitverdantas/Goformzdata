import re
import os
import openpyxl

# ----------------------------------------------------------------------
# CONFIG
# ----------------------------------------------------------------------
INPUT_FILE  = "PavanWorkDownloadedSheet.xlsx"   # your workbook
SHEET_NAME  = "Combined Sheet"                  # tab name
OUTPUT_FILE = "PavanWorkDownloadedSheet_ReportType.xlsx"

FORMNAME_COL = 2   # column B (Form Name)
NEWCOL_POS   = 3   # insert new column at position C (after B)


# ----------------------------------------------------------------------
# 1) EXTRACT report type (mirrors your REGEXEXTRACT logic)
#    Start : after the 2nd underscore
#    End   : before IR# | ID | SOL. | COPY | FINAL | end-of-string
# ----------------------------------------------------------------------
PATTERN = re.compile(
    r"^(?:[^_]*_){2}(.*?)"
    r"(?=\s*_?IR#|\s+ID\s|\s+SOL\.|\s+COPY|\s+FINAL|$)",
    re.IGNORECASE,
)


def excel_proper(text: str) -> str:
    """Capitalize first letter of each word, like Excel PROPER()."""
    return re.sub(r"[A-Za-z]+", lambda m: m.group(0).capitalize(), text)


def extract_report_type(form_name) -> str:
    if not form_name or not isinstance(form_name, str):
        return ""

    m = PATTERN.search(form_name)
    if not m:
        return ""

    result = m.group(1).strip()          # TRIM
    result = excel_proper(result)        # PROPER

    # SUBSTITUTE fixes (same as your formula)
    result = result.replace(" Id ", " ID ")
    result = result.replace("(Field)", "(FIELD)")
    result = result.replace("(Fab Shop)", "(FAB SHOP)")

    return result


# ----------------------------------------------------------------------
# 2) LOAD, INSERT COLUMN, FILL, SAVE
# ----------------------------------------------------------------------
def main():
    print("Loading workbook...")

    if not os.path.exists(INPUT_FILE):
        raise SystemExit(f"File not found: {INPUT_FILE}  (put it next to this script)")

    try:
        wb = openpyxl.load_workbook(INPUT_FILE)
    except PermissionError:
        raise SystemExit(f"'{INPUT_FILE}' is OPEN in Excel. Close it and re-run.")

    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Sheet '{SHEET_NAME}' not found. Tabs available: {wb.sheetnames}")

    ws = wb[SHEET_NAME]

    print("Inserting 'Report Type' column after column B...")
    ws.insert_cols(NEWCOL_POS)                              # new empty column at C
    ws.cell(row=1, column=NEWCOL_POS, value="Report Type")  # header

    print("Filling values...")
    count = 0
    for row in range(2, ws.max_row + 1):
        form_name = ws.cell(row=row, column=FORMNAME_COL).value
        ws.cell(row=row, column=NEWCOL_POS, value=extract_report_type(form_name))
        count += 1

    print(f"Saving... ({count} rows processed)")
    try:
        wb.save(OUTPUT_FILE)
    except PermissionError:
        raise SystemExit(f"'{OUTPUT_FILE}' is OPEN in Excel. Close it and re-run.")

    print(f"Done! Saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
